"""
Anonimiza a coluna 'Descrição' de relatórios de Ouvidoria (formato TCESP)
substituindo dados pessoais por marcadores tipados ([NOME], [CPF], ...).

Pipeline híbrido (todas as etapas rodam 100% local):
  1) Regex BR (CPF, CNPJ, RG, CNS, telefone, e-mail, CEP, placa, data)
  2) Heurísticas contextuais (Nome:/Mãe:/Pai:/Nasc:/Paciente:/Sr./Sra./Dr.)
  3) Endereços (Rua/Av./Estrada + texto + número)
  4) NER spaCy via Presidio (PERSON apenas — LOCATION é tratado pelas heurísticas)
  5) Substituição por marcadores tipados, preservando offsets corretos
  6) Geração de XLSX final + relatório de cobertura

Uso:
  python scripts/anonimizar_relatorio_tcesp.py \\
      --input  Item_44_3_Dados_Ouv_2025_ComDescricao_Mogi_das_Cruzes_SP_TCESP.xlsx \\
      --output Item_44_3_Dados_Ouv_2025_Anonimizado_Mogi_das_Cruzes_SP_TCESP.xlsx \\
      --coluna "Descrição" \\
      [--amostra 50]   # processa só N linhas (para validar)
      [--auditoria]    # gera coluna extra com texto original ao lado
"""

import argparse
import re
import sys
import time
from collections import Counter, defaultdict
from copy import copy

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from presidio_analyzer import AnalyzerEngine
from presidio_analyzer.nlp_engine import NlpEngineProvider



# 1) REGEX BR — alta precisão, rodam ANTES do NER

REGEX_BR = [
    ("CNPJ",     re.compile(r"\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b")),
    ("CNPJ",     re.compile(r"(?i)\bCNPJ[:\s]+(\d{14})\b")),
    ("CPF",      re.compile(r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b")),
    ("CPF",      re.compile(r"(?i)\bCPF[:\s]+(\d{11})\b")),
    ("RG",       re.compile(r"\b\d{2}\.\d{3}\.\d{3}-[\dxX]\b")),
    ("RG",       re.compile(r"(?i)\bRG[:\s]+(\d{7,10}[\dxX]?)\b")),
    ("CNS",      re.compile(r"(?i)\b(?:CNS|cart[aã]o\s+sus)[:\s]+\d{15}\b")),
    ("CNS",      re.compile(r"\b\d{3}\s\d{4}\s\d{4}\s\d{4}\b")),
    ("DOCUMENTO", re.compile(r"(?i)\bSIS\s*[:\-]?\s*\d{3,15}\b")),
    ("DOCUMENTO", re.compile(r"(?i)\bprontu[aá]rio\s*[:\-]?\s*\d{3,15}\b")),
    ("DOCUMENTO", re.compile(r"(?i)\bmatr[ií]cula\s*[:\-]?\s*\d{3,15}\b")),
    ("EMAIL",    re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b")),
    ("CEP",      re.compile(r"\b\d{5}-\d{3}\b")),
    ("PLACA",    re.compile(r"\b[A-Z]{3}[-\s]?\d[A-Z0-9]\d{2}\b")),
    ("TELEFONE", re.compile(r"\(?\d{2}\)?[\s\-]?9?\d{4}[\s\-]?\d{4}\b")),
    ("DATA",     re.compile(r"\b\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4}\b")),
]



# 2) HEURÍSTICAS CONTEXTUAIS — padrões típicos de ouvidoria

HEURISTICAS = [
    ("DATA", re.compile(
        r"(?i:\b(?:nasc(?:imento)?|dt\.?\s*nasc|data\s+de\s+nasc(?:imento)?)\b)"
        r"\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"
    )),
]

ANCORAS_FORTES_RX = re.compile(
    r"(?i)\b("
    r"nome\s+d[ao]\s+m[aã]e|nome\s+d[ao]\s+pai|"
    r"nome|paciente|"
    r"m[aã]e|pai|filho|filha|"
    r"sr\.?|sra\.?|dr\.?|dra\.?|prof\.?|profa\.?"
    r")\b\s*[:\-]?\s+"
)

ANCORAS_FRACAS_RX = re.compile(
    r"(?i)\b("
    r"requerente|solicitante|reclamante|denunciado|autor|respons[aá]vel"
    r")\b\s*[:\-]\s+"
)



# 3) ENDEREÇOS — rua/av/estrada + texto + número (não captura "rua" isolada)

REGEX_ENDERECO = re.compile(
    r"(?i)\b(?:rua|r\.|av\.?|avenida|alameda|al\.?|travessa|tv\.?|"
    r"estrada|estr\.?|rod\.?|rodovia|pra[çc]a|largo|servid[aã]o|caminho)\s+"
    r"[A-Za-zÀ-ÿ\s\-]{3,60}?"
    r",?\s*n?[°º]?\s*\d{1,6}\b"
)



# 3.b) STOPWORDS DE DOMÍNIO — palavras que NUNCA são PII (mesmo capitalizadas)

STOPWORDS_DOMINIO = {
    # Atores de ouvidoria
    "requerente", "solicitante", "reclamante", "denunciante", "autor",
    "munícipe", "municipe", "usuário", "usuario", "paciente", "responsável",
    "responsavel", "interessado", "cidadão", "cidadao",
    # Verbos de relato
    "reclama", "reclamou", "reclamar", "reclamando", "reclam",
    "solicita", "solicitou", "solicitar", "solicitando",
    "denuncia", "denunciou", "informa", "informou",
    "relata", "relatou", "relatando", "relação", "relacao",
    "pede", "pediu", "pedir", "pedindo",
    "aguarda", "aguardou", "aguardava", "aguardando", "aguardar",
    "correndo", "correu", "corre", "comparecer",
    "compareceu", "comparecend", "compareceram", "corrrendo",
    "ocorrendo", "ocorre", "ocorreu", "ocorrência", "ocorrencia",
    "diz", "disse", "dizendo", "dizer",
    "deixou", "deixar", "deixando", "deixa", "deixaram",
    "chegou", "chega", "chegando", "saiu", "sai", "saindo",
    "tem", "teve", "tendo", "ter", "havia", "haver",
    "foi", "vai", "indo", "veio", "vem", "vindo",
    "estão", "estao", "está", "esta", "estava", "estive",
    "há", "ha", "havia", "houve", "uma", "um", "umas", "uns",
    "elogia", "elogiou", "elogiando", "elogiar",
    "agradeço", "agradeces", "agradece", "agradeceu",
    "atendimento", "atendimentos", "atende", "atendeu", "atender",
    "serviço", "servico", "serviços", "servicos", "trabalho",
    "morador", "moradora", "moradores", "moradoras",
    "trabalhador", "trabalhadora", "agente", "fiscal", "professor",
    "professora", "professores", "professoras", "profissional",
    # Relações
    "esposa", "esposo", "marido", "sogra", "sogro", "neto", "neta",
    "filho", "filha", "filhos", "filhas", "irmão", "irmao", "irmã",
    "primo", "prima", "tio", "tia", "avô", "avo", "avó",
    "vizinho", "vizinha", "vizinhos",
    # Rótulos
    "nasc", "nascimento", "nasceu", "nascida", "nascido",
    "data", "datas", "ponto", "referência", "referencia", "contato",
    "telefone", "fone", "celular", "endereço", "endereco", "rua",
    "avenida", "estrada", "bairro", "cidade", "estado",
    "nome", "mãe", "mae", "pai", "guarda", "protocolo",
    # Pronomes e auxiliares comuns que viram "Title" em caixa alta
    "sua", "seu", "este", "esta", "esse", "essa", "aquele", "aquela",
    "outro", "outra", "outros", "outras", "muito", "muita",
    "grandes", "grande", "pequeno", "pequena",
    "diversas", "diversos", "vezes", "muitas", "muitos",
    "última", "ultima", "último", "ultimo",
    # Meses e advérbios temporais (frequentes em textos de ouvidoria)
    "janeiro", "fevereiro", "março", "marco", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
    "hoje", "ontem", "amanhã", "amanha", "desde", "ainda", "sempre",
    "nunca", "agora", "logo", "antes", "depois", "também", "tambem",
    "porém", "porem", "assim", "então", "entao", "porque",
    "que", "quando", "quanto", "como", "onde", "neto", "neta", "netos",
    # Termos administrativos
    "ubs", "upa", "caps", "semae", "secretaria", "ouvidoria",
    "câmara", "camara", "prefeitura", "civil", "guarda",
    # Risco de aparecer só em CAIXA ALTA
    "risco", "riscos", "providências", "providencias", "vistoria",
    "imóvel", "imovel", "residência", "residencia", "obra",
    "carro", "moto", "ônibus", "onibus", "veículo", "veiculo",
    "som", "carga", "barulho",
}

# Conectores só são removidos das bordas quando estiverem entre stopwords/conectores.
CONECTORES = {"da", "de", "do", "das", "dos", "e", "na", "no", "nas", "nos",
              "com", "sem", "para", "por", "à", "às", "ao", "aos"}


def _norm_word(s: str) -> str:
    return s.lower().strip(".,;:()[]\"'-!?")


# Verbos de relato comuns — usados também na expansão de nomes por âncora
VERBOS_INTERNOS_NOME = {
    "reclama", "reclamou", "reclamando", "solicita", "solicitou",
    "denuncia", "denunciou", "informa", "informou", "relata", "relatou",
    "aguarda", "aguardou", "aguardando", "aguardava",
    "pede", "pediu", "diz", "disse", "deixou", "tem", "teve",
    "chegou", "veio", "vai", "foi", "está", "estão", "estava",
    "ocorrendo", "ocorre", "ocorreu", "comparecend", "compareceu",
    "que", "porque", "quando", "como", "onde", "neto", "neta",
}


def detectar_nomes_por_ancora(texto: str):
    """
    Estratégia robusta para detecção de nomes próprios após rótulos comuns.
    Para cada ocorrência de âncora (Nome:, Paciente:, Sr., etc.), expande
    token-a-token capturando palavras capitalizadas até encontrar:
      - uma stopword/verbo
      - uma palavra não-capitalizada (que não seja conector entre nomes)
      - fim de texto / pontuação forte

    Vantagem sobre regex única: detecta TODAS as âncoras (a primeira não
    consome o texto), funciona em CAIXA ALTA, e tem trim incorporado.
    """
    spans = []
    TOKEN_RX = re.compile(r"\S+")

    todas_ancoras = []
    for am in ANCORAS_FORTES_RX.finditer(texto):
        todas_ancoras.append((am.end(), False)) 
    for am in ANCORAS_FRACAS_RX.finditer(texto):
        todas_ancoras.append((am.end(), True))

    for start, exige_2_tokens in todas_ancoras:
        sub = texto[start:start + 200]
        toks = list(TOKEN_RX.finditer(sub))
        if not toks:
            continue

        coletados = []
        for idx, tm in enumerate(toks):
            t = tm.group(0)
            limpo = t.strip(".,;:()[]\"'-!?")
            wnorm = _norm_word(t)

            if any(p in t for p in [".", "!", "?", ";"]) and idx > 0:
                break

            if ":" in t:
                break

            if wnorm in STOPWORDS_DOMINIO or wnorm in VERBOS_INTERNOS_NOME:
                break

            if wnorm in CONECTORES:
                if coletados and idx + 1 < len(toks):
                    nxt = toks[idx + 1].group(0).lstrip(".,;:()[]\"'-!?")
                    if nxt and nxt[0].isalpha() and nxt[0].isupper():
                        nxt_norm = _norm_word(nxt)
                        if (nxt_norm not in STOPWORDS_DOMINIO and
                                nxt_norm not in VERBOS_INTERNOS_NOME):
                            coletados.append(tm)
                            continue
                break

            if limpo and limpo[0].isalpha() and limpo[0].isupper():
                coletados.append(tm)
                puros = sum(1 for c in coletados
                            if _norm_word(c.group(0)) not in CONECTORES)
                if puros >= 6:
                    break
                continue

            break

        if not coletados:
            continue

        puros_coletados = [c for c in coletados
                           if _norm_word(c.group(0)) not in CONECTORES]
        if not puros_coletados:
            continue

        if exige_2_tokens and len(puros_coletados) < 2:
            continue

        if (not exige_2_tokens and len(puros_coletados) == 1 and
                len(puros_coletados[0].group(0).strip(".,;:()[]\"'-!?")) < 4):
            continue

        ns = start + coletados[0].start()
        ne = start + coletados[-1].end()
        spans.append((ns, ne, "NOME"))

    return spans


def parece_nome_proprio(start: int, end: int, texto: str) -> bool:
    """
    Verifica se um span efetivamente parece um nome próprio (vs. uma frase
    comum em CAIXA ALTA que o NER confundiu). Critérios:
      - Pelo menos 1 token "puro" (não stopword, não conector)
      - >= 60% das palavras puras começam com letra maiúscula no texto original
      - Nenhuma palavra interna é um verbo de relato típico de ouvidoria
    """
    sub = texto[start:end]
    tokens = re.findall(r"\S+", sub)
    if not tokens:
        return False

    puros = [t for t in tokens
             if _norm_word(t) not in STOPWORDS_DOMINIO
             and _norm_word(t) not in CONECTORES]
    if not puros:
        return False

    VERBOS_INTERNOS = {
        "reclama", "reclamou", "solicita", "solicitou", "denuncia",
        "informa", "informou", "relata", "relatou", "aguarda", "aguardou",
        "pede", "pediu", "diz", "disse", "deixou", "tem", "teve",
        "chegou", "veio", "vai", "foi", "está", "estão", "estava",
        "ocorrendo", "comparecend", "compareceu",
    }
    for t in tokens:
        if _norm_word(t) in VERBOS_INTERNOS:
            return False

    capitalizados = 0
    for t in puros:
        primeiro = t.lstrip(".,;:()[]\"'-!?")
        if primeiro and primeiro[0].isupper():
            capitalizados += 1
    return (capitalizados / len(puros)) >= 0.6


def trim_span_stopwords(start: int, end: int, texto: str):
    """
    Recorta um span removendo stopwords e conectores das bordas. Conectores
    só são cortados se a palavra adjacente também for stopword/conector
    (evita cortar nomes legítimos como "Maria de Lourdes").

    Retorna (novo_start, novo_end) ou (None, None) se sobrar vazio/curto.
    """
    sub = texto[start:end]
    tokens = list(re.finditer(r"\S+", sub))
    if not tokens:
        return None, None

    def is_stop(idx):
        return _norm_word(tokens[idx].group(0)) in STOPWORDS_DOMINIO

    def is_conn(idx):
        return _norm_word(tokens[idx].group(0)) in CONECTORES

    i, j = 0, len(tokens) - 1

    while i <= j:
        if is_stop(i):
            i += 1
        elif is_conn(i) and i + 1 <= j and (is_stop(i + 1) or is_conn(i + 1)):
            i += 1
        else:
            break

    while j >= i:
        if is_stop(j):
            j -= 1
        elif is_conn(j) and j - 1 >= i and (is_stop(j - 1) or is_conn(j - 1)):
            j -= 1
        else:
            break

    if i > j:
        return None, None

    novo_start = start + tokens[i].start()
    novo_end = start + tokens[j].end()
    if novo_end - novo_start < 3:
        return None, None
    return novo_start, novo_end


# 4) NORMALIZAÇÃO DE CAIXA — corrige texto em CAIXA ALTA para o NER funcionar

def normalizar_caixa(texto: str) -> str:
    """
    Se mais de 60% das letras do texto forem maiúsculas, converte palavras
    LONGAS (>= 5 letras) para Title Case. Palavras curtas (QUE, NO, DA, DIZ,
    FOI, etc.) ficam em MAIÚSCULAS para que o NER do spaCy NÃO as confunda
    com nomes próprios — esse é o principal efeito colateral da normalização.
    NÃO altera o texto original — só é usado internamente para o NER.
    O comprimento de cada char é preservado para manter os offsets corretos.
    """
    letras = [c for c in texto if c.isalpha()]
    if not letras:
        return texto
    pct_maiuscula = sum(1 for c in letras if c.isupper()) / len(letras)
    if pct_maiuscula < 0.6:
        return texto

    def title_word(m):
        w = m.group(0)
        if len(w) < 5:
            return w 
        return w[0] + w[1:].lower()

    return re.sub(r"[A-ZÁÉÍÓÚÂÊÔÃÕÇ]{2,}", title_word, texto)


# 5) ANALYZER PRESIDIO — só PERSON (LOCATION e DATE_TIME tratamos com regex)

def criar_analyzer():
    nlp_config = {
        "nlp_engine_name": "spacy",
        "models": [{"lang_code": "pt", "model_name": "pt_core_news_lg"}],
    }
    provider = NlpEngineProvider(nlp_configuration=nlp_config)
    nlp_engine = provider.create_engine()
    return AnalyzerEngine(nlp_engine=nlp_engine, supported_languages=["pt"])


# 6) MERGE DE SPANS — resolve sobreposições mantendo o de maior prioridade

PRIORIDADE = {
    "CPF": 10, "CNPJ": 10, "RG": 9, "CNS": 9, "EMAIL": 8, "CEP": 7,
    "TELEFONE": 7, "PLACA": 6, "DATA": 6, "ENDERECO": 5, "NOME": 4,
}


def merge_spans(spans):
    """
    Recebe lista de (start, end, tipo). Resolve sobreposições mantendo o de
    maior prioridade; em caso de empate, o maior comprimento; depois o primeiro.
    Retorna lista ordenada, sem overlaps.
    """
    if not spans:
        return []
    spans = sorted(spans, key=lambda s: (s[0], -(s[1] - s[0])))
    resultado = []
    for s in spans:
        if not resultado:
            resultado.append(s)
            continue
        last = resultado[-1]
        if s[0] >= last[1]:
            resultado.append(s)
        else:
            pri_s = PRIORIDADE.get(s[2], 0)
            pri_l = PRIORIDADE.get(last[2], 0)
            comp_s = s[1] - s[0]
            comp_l = last[1] - last[0]
            if (pri_s, comp_s) > (pri_l, comp_l):
                resultado[-1] = s
    return resultado


# 7) PIPELINE PRINCIPAL DE ANONIMIZAÇÃO DE UMA DESCRIÇÃO

def anonimizar(texto: str, analyzer, contador: Counter) -> str:
    if not texto or not isinstance(texto, str):
        return texto

    spans = []

    for tipo, rx in REGEX_BR:
        for m in rx.finditer(texto):
            spans.append((m.start(), m.end(), tipo))

    for tipo, rx in HEURISTICAS:
        for m in rx.finditer(texto):
            if m.groups():
                spans.append((m.start(1), m.end(1), tipo))
            else:
                spans.append((m.start(), m.end(), tipo))

    for ns, ne, tipo in detectar_nomes_por_ancora(texto):
        spans.append((ns, ne, tipo))

    for m in REGEX_ENDERECO.finditer(texto):
        spans.append((m.start(), m.end(), "ENDERECO"))

    spans_existentes = [(s[0], s[1]) for s in spans]

    def sobrepoe(start, end):
        for es, ee in spans_existentes:
            if start < ee and es < end:
                return True
        return False

    texto_norm = normalizar_caixa(texto)
    try:
        resultados = analyzer.analyze(
            text=texto_norm,
            language="pt",
            entities=["PERSON"],
            score_threshold=0.5,
        )
        for r in resultados:
            if sobrepoe(r.start, r.end):
                continue
            ns, ne = trim_span_stopwords(r.start, r.end, texto)
            if ns is None:
                continue
            if not parece_nome_proprio(ns, ne, texto):
                continue
            spans.append((ns, ne, "NOME"))
    except Exception as e:
        print(f"  [warn] NER falhou em descrição: {e}", file=sys.stderr)

    spans = merge_spans(spans)

    if not spans:
        return texto

    resultado = texto
    for start, end, tipo in sorted(spans, key=lambda s: -s[0]):
        marcador = f"[{tipo}]"
        resultado = resultado[:start] + marcador + resultado[end:]
        contador[tipo] += 1

    return resultado


# 8) IO XLSX — preserva todas as colunas, atualiza só a coluna alvo

def processar_xlsx(input_path, output_path, coluna_alvo, amostra=None,
                   auditoria=False):
    print(f"[1/4] Abrindo {input_path}...")
    wb_in = load_workbook(input_path, read_only=False, data_only=True)
    ws_in = wb_in.active

    header = [c.value for c in ws_in[1]]
    if coluna_alvo not in header:
        print(f"ERRO: coluna {coluna_alvo!r} não encontrada. Colunas: {header}")
        sys.exit(1)
    col_idx = header.index(coluna_alvo)
    col_letter = ws_in.cell(row=1, column=col_idx + 1).column_letter

    total_linhas = ws_in.max_row - 1
    if amostra:
        total_linhas = min(total_linhas, amostra)

    print(f"[2/4] Carregando modelo NER (pode demorar ~10s na 1ª vez)...")
    t0 = time.time()
    analyzer = criar_analyzer()
    print(f"      Modelo pronto em {time.time()-t0:.1f}s")

    print(f"[3/4] Anonimizando {total_linhas} linhas...")
    contador = Counter()
    afetadas = 0
    t_start = time.time()

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = ws_in.title

    new_header = list(header)
    if auditoria:
        new_header.append(f"{coluna_alvo} (original)")
    ws_out.append(new_header)

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    for cell in ws_out[1]:
        cell.font = header_font
        cell.fill = header_fill

    rows_iter = ws_in.iter_rows(min_row=2, values_only=True)
    for i, row in enumerate(rows_iter):
        if amostra and i >= amostra:
            break

        row = list(row)
        original = row[col_idx]
        if isinstance(original, str) and original.strip():
            antes = sum(contador.values())
            anonimizado = anonimizar(original, analyzer, contador)
            depois = sum(contador.values())
            if depois > antes:
                afetadas += 1
            row[col_idx] = anonimizado

        if auditoria:
            row.append(original)

        ws_out.append(row)

        if (i + 1) % 500 == 0:
            elapsed = time.time() - t_start
            rate = (i + 1) / elapsed
            eta = (total_linhas - (i + 1)) / rate if rate > 0 else 0
            print(f"      {i+1:>6}/{total_linhas} ({100*(i+1)/total_linhas:5.1f}%) "
                  f"| {rate:5.1f} linhas/s | ETA {eta/60:5.1f} min")

    elapsed = time.time() - t_start
    print(f"      Concluído em {elapsed/60:.1f} min ({total_linhas/elapsed:.1f} linhas/s)")

    print(f"[4/4] Ajustando largura/wrap da coluna {col_letter} e salvando...")
    ws_out.column_dimensions[col_letter].width = 80
    for r in range(2, ws_out.max_row + 1):
        ws_out.cell(row=r, column=col_idx + 1).alignment = Alignment(
            wrap_text=True, vertical="top"
        )

    wb_out.save(output_path)
    print(f"      Salvo em: {output_path}")

    return contador, afetadas, total_linhas


# 9) RELATÓRIO DE COBERTURA

def imprimir_relatorio(contador, afetadas, total):
    print("\n" + "=" * 60)
    print("RELATÓRIO DE COBERTURA DE ANONIMIZAÇÃO")
    print("=" * 60)
    print(f"Linhas processadas:                {total}")
    print(f"Linhas com PII detectada:          {afetadas} "
          f"({100*afetadas/total:.1f}%)")
    print(f"Total de entidades mascaradas:     {sum(contador.values())}")
    print()
    print("Detalhamento por tipo:")
    print(f"  {'TIPO':<12} {'OCORRÊNCIAS':>12}")
    print(f"  {'-'*12} {'-'*12}")
    for tipo, n in contador.most_common():
        print(f"  {tipo:<12} {n:>12}")
    print("=" * 60)

def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", required=True, help="Arquivo XLSX de entrada.")
    parser.add_argument("--output", required=True, help="Arquivo XLSX de saída.")
    parser.add_argument("--coluna", default="Descrição",
                        help="Nome da coluna a anonimizar (padrão: 'Descrição').")
    parser.add_argument("--amostra", type=int, default=None,
                        help="Processar apenas N linhas (para teste).")
    parser.add_argument("--auditoria", action="store_true",
                        help="Adicionar coluna extra com o texto original.")
    args = parser.parse_args()

    contador, afetadas, total = processar_xlsx(
        args.input, args.output, args.coluna,
        amostra=args.amostra, auditoria=args.auditoria,
    )
    imprimir_relatorio(contador, afetadas, total)


if __name__ == "__main__":
    main()
